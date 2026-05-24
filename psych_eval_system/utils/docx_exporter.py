"""
Exports the final markdown report to a .docx file.

PII is re-injected here — this is the only place real patient identifiers
touch the output. The LLM-generated markdown uses tokens throughout.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Optional

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from psych_eval_system.models.patient import PatientPII
from psych_eval_system.models.report import FullReport
from psych_eval_system.utils.pii_sanitizer import PIISanitizer


class DocxExporter:
    def __init__(self, patient: PatientPII, sanitizer: PIISanitizer) -> None:
        self._patient = patient
        self._sanitizer = sanitizer

    def export(self, report: FullReport, output_path: Path) -> Path:
        doc = Document()
        self._set_default_style(doc)
        self._add_header(doc)

        markdown = self._sanitizer.restore(report.final_markdown or report.draft_markdown)
        self._render_markdown(doc, markdown)

        self._add_footer(doc)
        doc.save(output_path)
        return output_path

    def _set_default_style(self, doc: Document) -> None:
        style = doc.styles["Normal"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)

    def _add_header(self, doc: Document) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("CONFIDENTIAL\nPSYCHOLOGICAL EVALUATION")
        run.bold = True
        run.font.size = Pt(14)

        doc.add_paragraph()  # spacer

        meta_lines = [
            f"Client Name: {self._patient.full_name}",
            f"Birth Date: {self._patient.birth_date.strftime('%m/%d/%Y')}",
            f"Age: {self._patient.age_years}-years, {self._patient.age_months}-months",
            f"Sex: {self._patient.gender}",
            f"Evaluation Date: {self._patient.referral_info.evaluation_date or ''}",
            f"Date of Report: {self._patient.referral_info.report_date or ''}",
        ]
        for line in meta_lines:
            p = doc.add_paragraph(line)
            p.runs[0].font.size = Pt(12)

        doc.add_paragraph()

    def _add_footer(self, doc: Document) -> None:
        doc.add_paragraph()
        p = doc.add_paragraph(
            f"This report was prepared by {self._patient.referral_info.examiner or '[Examiner]'} "
            f"at {self._patient.referral_info.clinic_name}. "
            "All findings are confidential and intended for the named recipient only. "
            "A licensed psychologist must review this document before clinical use."
        )
        p.runs[0].font.size = Pt(10)
        p.runs[0].italic = True

    def _render_markdown(self, doc: Document, markdown: str) -> None:
        for line in markdown.split("\n"):
            stripped = line.strip()
            if not stripped:
                doc.add_paragraph()
                continue

            if stripped.startswith("### "):
                p = doc.add_heading(stripped[4:], level=3)
            elif stripped.startswith("## "):
                p = doc.add_heading(stripped[3:], level=2)
            elif stripped.startswith("# "):
                p = doc.add_heading(stripped[2:], level=1)
            elif stripped.startswith("**") and stripped.endswith("**"):
                p = doc.add_paragraph()
                run = p.add_run(stripped.strip("*"))
                run.bold = True
            elif stripped.startswith("|"):
                # Minimal table rendering — multi-line tables handled as paragraphs
                p = doc.add_paragraph(stripped)
                p.runs[0].font.name = "Courier New"
                p.runs[0].font.size = Pt(9)
            else:
                p = doc.add_paragraph(stripped)
                # Inline bold: **text**
                if "**" in stripped:
                    p.clear()
                    parts = re.split(r"\*\*", stripped)
                    for i, part in enumerate(parts):
                        if part:
                            run = p.add_run(part)
                            run.bold = (i % 2 == 1)
